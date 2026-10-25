#!/usr/bin/env python3
"""准备指南分析语料：解压 → 抽取文本 → 判定地区 → 输出 docs.json。

原始 PDF **不进入本仓库**。默认把语料放在 ~/.hermes/cache/scratch/guidelines，
归档副本放 efgartigimod-wiki raw/papers/（见 report/runbooks/monthlyDivergenceAnalysis.md）。

用法：
    python3 scripts/prepare-guideline-corpus.py --zip ~/Downloads/指南共识.zip
    python3 scripts/prepare-guideline-corpus.py --dir ~/Downloads/指南共识 --scratch /tmp/gl
"""
from __future__ import annotations

import argparse
import json
import os
import re
import zipfile
from pathlib import Path

DEFAULT_SCRATCH = os.environ.get(
    'MG_GUIDELINE_SCRATCH', str(Path.home() / '.hermes/cache/scratch/guidelines'))

# 用于判定"中国指南/共识"的机构与刊名线索（只看正文前 4000 字）
CN_ORG_HINTS = [
    '中华医学会', '中华神经科杂志', '中华眼底病杂志', '中华胸心血管外科杂志', '中华护理杂志',
    '中国神经免疫', '中国免疫学会', '中国罕见病', '中国药房', '中国胸心血管外科临床杂志',
    '神经免疫学组', '北京神经科学学会', '中国抗癌协会', '中国重症肌无力联盟',
]
# 文件名本身不足以判定的特例（正文是中文但文件名不含中文机构名）
REGION_OVERRIDES: dict[str, str] = {}


def fix_zip_name(raw: str) -> str:
    """zip 内 GBK 文件名字节会被 zipfile 解成 cp437，需回转。"""
    try:
        return raw.encode('cp437').decode('gbk')
    except Exception:
        try:
            return raw.encode('cp437').decode('utf-8')
        except Exception:
            return raw


def extract_pdf(path: Path) -> str:
    import fitz
    doc = fitz.open(path)
    try:
        return '\n'.join(page.get_text('text') for page in doc)
    finally:
        doc.close()


def extract_docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(' | '.join(cell.text.strip() for cell in row.cells))
    return '\n'.join(parts)


def slice_mg_chapter(text: str, budget: int = 60000) -> str:
    """多病种合集（如中国罕见病诊疗指南）只取重症肌无力章节。"""
    hits = [m.start() for m in re.finditer('重症肌无力', text)]
    if not hits:
        return text[:budget]
    head = hits[0]
    return text[max(0, head - 200):head + budget]


def detect_region(name: str, text: str) -> str:
    if '台湾' in name:
        return 'taiwan'
    head = text[:4000]
    zh = len(re.findall(r'[\u4e00-\u9fff]', head))
    ratio = zh / max(1, len(head))
    if any(hint in head for hint in CN_ORG_HINTS) or ratio > 0.3:
        return 'china'
    return 'international'


def unzip_all(zip_path: Path, out_dir: Path) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            name = fix_zip_name(info.filename)
            safe = re.sub(r'[\\/:*?"<>|]', '_', name).strip()
            dest = out_dir / safe
            with zf.open(info) as src, open(dest, 'wb') as dst:
                dst.write(src.read())
            paths.append(dest)
    return paths


def main() -> int:
    ap = argparse.ArgumentParser(description='准备中外指南差异分析语料')
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument('--zip', help='指南压缩包（GBK 文件名可用）')
    src.add_argument('--dir', help='已解压的指南目录')
    ap.add_argument('--scratch', default=DEFAULT_SCRATCH, help='语料输出目录')
    args = ap.parse_args()

    scratch = Path(args.scratch).expanduser()
    pdf_dir = scratch / 'pdf'
    text_dir = scratch / 'texts'
    text_dir.mkdir(parents=True, exist_ok=True)

    if args.zip:
        files = unzip_all(Path(args.zip).expanduser(), pdf_dir)
    else:
        files = sorted(p for p in Path(args.dir).expanduser().iterdir() if p.is_file())

    rows = []
    for i, path in enumerate(files):
        suffix = path.suffix.lower()
        if suffix == '.pdf':
            raw = extract_pdf(path)
        elif suffix == '.docx':
            raw = extract_docx(path)
        else:
            continue
        # 多病种合集（如中国罕见病诊疗指南，>20 万字符）只保留 MG 章节，
        # 否则一份文件会把整本指南塞进立场提取，既贵又跑偏。
        sliced = False
        if len(raw) > 200_000 and '重症肌无力' in raw:
            raw = slice_mg_chapter(raw)
            sliced = True
        text = re.sub(r'\n{3,}', '\n\n', raw)
        text = re.sub(r'[ \t]{2,}', ' ', text)
        key = f'g{i + 1:02d}'
        (text_dir / f'{key}.txt').write_text(text, encoding='utf-8')
        year = re.search(r'(19|20)\d{2}', path.name)
        rows.append({
            'key': key,
            'name': path.name,
            'year': year.group(0) if year else '',
            'region': REGION_OVERRIDES.get(key) or detect_region(path.name, text),
            'lang': 'zh' if len(re.findall(r'[\u4e00-\u9fff]', text)) > 200 else 'en',
            'chars': len(text),
            'chapter_sliced': sliced,
            'path': str(path),
        })

    (scratch / 'docs.json').write_text(
        json.dumps(rows, ensure_ascii=False, indent=1), encoding='utf-8')

    counts: dict[str, int] = {}
    for row in rows:
        counts[row['region']] = counts.get(row['region'], 0) + 1
    print(f'语料目录：{scratch}')
    print(f'抽取 {len(rows)} 份：' + ' / '.join(f'{k} {v}' for k, v in sorted(counts.items())))
    print('下一步：python3 scripts/guideline-stance-extract.py --scratch ' + str(scratch))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
